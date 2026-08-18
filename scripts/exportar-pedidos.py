#!/usr/bin/env python3
"""Exporta el checklist de pedidos a JSON, para publicarlo en el sitio de progreso.

El seguimiento vive en docs/pedidos-alan.xlsx porque ahí es cómodo de editar
con fórmulas y de mandar por chat. Pero un .xlsx en GitHub Pages no se lee: el
navegador lo ofrece para descargar y ahí se queda, así que el sitio no puede
mostrarlo directamente.

Este script corre a mano, no en cada build:

    python3 scripts/exportar-pedidos.py

Lee el Excel y escribe docs/pedidos-alan.json, que es lo que
scripts/construir-sitio.mjs de verdad lee para armar la página. El JSON queda
commiteado junto con el Excel: cada vez que se edita uno hay que correr esto y
commitear los dos juntos. Es el mismo patrón que scripts/planisferio.mjs.

No usa ninguna librería de Node para leer Excel —evita sumar una dependencia
sólo para esto, y de paso evita las que tienen CVEs conocidos sin parchear en
npm— porque este archivo ya se edita con openpyxl durante toda la sesión.
"""

import json
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "docs" / "pedidos-alan.xlsx"
DESTINO = RAIZ / "docs" / "pedidos-alan.json"


def principal() -> None:
    wb = openpyxl.load_workbook(ORIGEN, data_only=True)
    ws = wb["Pedidos"]

    filas = []
    fila = 5
    while ws.cell(row=fila, column=3).value:
        filas.append(
            {
                "numero": ws.cell(row=fila, column=1).value,
                "area": ws.cell(row=fila, column=2).value,
                "pedido": ws.cell(row=fila, column=3).value,
                "estado": ws.cell(row=fila, column=4).value,
                "observaciones": ws.cell(row=fila, column=5).value or "",
                "prioridad": ws.cell(row=fila, column=6).value,
            }
        )
        fila += 1

    DESTINO.write_text(
        json.dumps({"filas": filas}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"{len(filas)} filas exportadas a {DESTINO.relative_to(RAIZ)}")


if __name__ == "__main__":
    principal()
